#!/usr/bin/env python3
"""Build StoreVerse PRD as a Word document with architecture and flow diagrams."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DIAGRAM = ROOT / "docs" / "assets" / "diagrams"
OUT = ROOT / "docs" / "PRD-门店颗粒度一线赋能.docx"

INK = RGBColor(0x1B, 0x17, 0x14)
ACCENT = RGBColor(0xD2, 0x4B, 0x16)
MUTED = RGBColor(0x6D, 0x64, 0x5B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_BG = "1F1814"
ROW_BG = "F7F3EC"
FONT_EA = "Hiragino Sans GB"
FONT_LATIN = "Calibri"


def set_run_font(run, size=11, bold=False, color=INK, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = FONT_LATIN
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), FONT_LATIN)
    rfonts.set(qn("w:hAnsi"), FONT_LATIN)
    rfonts.set(qn("w:eastAsia"), FONT_EA)
    rfonts.set(qn("w:cs"), FONT_LATIN)


def shade_cell(cell, fill: str):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "C8BFB2")
        tc_borders.append(el)
    tc_pr.append(tc_borders)


def cell_text(cell, text, *, bold=False, size=10, color=INK, align="left"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    set_cell_border(cell)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, HEADER_BG)
        cell_text(cell, header, bold=True, size=10, color=WHITE, align="center")
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.rows[r + 1].cells[c]
            if r % 2 == 1:
                shade_cell(cell, ROW_BG)
            cell_text(cell, value, size=10)
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size={1: 16, 2: 13, 3: 12}.get(level, 12), bold=True, color=INK)
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    return p


def add_body(doc, text, *, bold=False, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, size=11, bold=bold)
    return p


def add_plain(doc, text, *, size=11, bold=False, color=INK, align=None, space_after=6):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    if p.runs:
        p.runs[0].text = text
        set_run_font(p.runs[0], size=11)
    else:
        run = p.add_run(text)
        set_run_font(run, size=11)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(text)
    set_run_font(run, size=9, italic=True, color=MUTED)
    return p


def add_figure(doc, filename, caption, width_cm=15.4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(DIAGRAM / filename), width=Cm(width_cm))
    add_caption(doc, caption)


def set_style_fonts(doc):
    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "Title"):
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        style.font.name = FONT_LATIN
        style.font.color.rgb = INK
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:ascii"), FONT_LATIN)
        rfonts.set(qn("w:hAnsi"), FONT_LATIN)
        rfonts.set(qn("w:eastAsia"), FONT_EA)


def add_page_number(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=9, color=MUTED)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)


def setup_header_footer(doc):
    section = doc.sections[0]
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("StoreVerse 门店运营 AI  ·  PRD  ·  v0.2")
    set_run_font(run, size=9, color=MUTED)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    left = fp.add_run("机密 · 仅供内部演示  ·  第 ")
    set_run_font(left, size=9, color=MUTED)
    add_page_number(fp)
    right = fp.add_run(" 页")
    set_run_font(right, size=9, color=MUTED)


def build():
    doc = Document()
    set_style_fonts(doc)

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    setup_header_footer(doc)

    add_plain(doc, "STOREVERSE", size=12, bold=True, color=ACCENT, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    add_plain(doc, "门店运营 AI 工具", size=28, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    add_plain(doc, "产品需求文档（PRD）v0.2", size=16, bold=True, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
    add_plain(
        doc,
        "DSR 拍照进店后先看销售看板，再生成 AI 运营建议；对每条建议选择采纳或忽略。总部只监控使用与采纳，不替代经营驾驶舱。",
        size=12,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=20,
    )
    add_table(
        doc,
        ["项", "内容"],
        [
            ["产品名称", "StoreVerse 门店运营 AI 工具（演示版）"],
            ["版本", "v0.2"],
            ["状态", "基于 2026-09-08 两轮面聊修订"],
            ["主路径", "DSR 看板 → 生成建议 → 采纳/忽略"],
            ["辅路径", "总部使用监控；导购为后续"],
            ["文档日期", "2026-09-08"],
        ],
    )

    add_heading(doc, "1. 一句话定义", 1)
    add_body(
        doc,
        "DSR 拍照识别门店后，先看该店销售看板（周边标签、各平台订单、Top 商品、机制效果），再一键生成 AI 运营建议；对每条建议选择采纳或忽略。总部后台只监控 DSR 用了什么、采纳了什么。",
    )

    add_heading(doc, "2. 相对 v0.1 的关键调整", 1)
    add_table(
        doc,
        ["原理解", "现结论"],
        [
            ["拍照后直接给做工指令", "先进入销售看板，建议由「生成运营建议」触发"],
            ["待办勾选 = 我去上架", "无法直连客户后台，改为采纳 / 忽略"],
            ["明确不做总部后台", "做轻量 DSR 使用监控（到店、报告、采纳）"],
            ["导购与 DSR 并列", "DSR 为主路径，导购后续"],
            ["到店/到家两块销量", "按微信支付、京东到家、美团闪购拆开，带订单和补贴"],
            ["建议直接罗列", "判断链：供给 → 动销 → 机制 → 补贴"],
            ["与客户门店系统拉通", "暂不拉通，门店是 StoreVerse 的"],
        ],
    )

    add_heading(doc, "3. 产品定位", 1)
    add_table(
        doc,
        ["维度", "定位"],
        [
            ["一线形态", "手机端。拍照识店，或从历史门店进入"],
            ["进店后第一屏", "该店销售看板，不是建议清单"],
            ["建议如何出现", "点击「生成运营建议」后由 AI 生成"],
            ["一线动作", "采纳 / 忽略，可上报总部做计划或配券"],
            ["总部形态", "监控 DSR 使用与建议采纳，不是全国报表"],
            ["门店主数据", "StoreVerse 门店，本期不与客户自有系统打通"],
        ],
    )

    add_heading(doc, "4. 能力架构", 1)
    add_body(doc, "一线先进看板，再生成建议；采纳结果回写总部监控。判断顺序固定为供给、动销、机制、补贴。")
    add_figure(doc, "01-capability-architecture.png", "图 1  产品能力架构：看板 → 生成建议 → 采纳 → 总部")
    add_figure(doc, "05-sequence.png", "图 2  端到端时序")
    add_figure(doc, "02-data-architecture.png", "图 3  数据架构：三平台订单、机制与周边标签")

    add_heading(doc, "5. DSR 主流程", 1)
    add_figure(doc, "03-dsr-flow.png", "图 4  DSR：进店看板 → 生成建议 → 采纳或忽略")
    add_bullet(doc, "新管或未管过的门：拍照识店。")
    add_bullet(doc, "管过的店：进入历史门店。")
    add_bullet(doc, "看板含：1–3km 标签，昨天/7天/30天的微信、京东到家、美团闪购订单与补贴，Top 商品，机制效果。")
    add_bullet(doc, "点击生成运营建议；AI 按供给→动销→机制→补贴输出。")
    add_bullet(doc, "每条建议采纳或忽略。采纳可上报总部制定计划或申请配券。")
    add_bullet(doc, "同一门店的历史生成记录可回看。")

    add_heading(doc, "6. AI 建议类型", 1)
    add_table(
        doc,
        ["类型", "示例", "判断链"],
        [
            ["供给", "缺货 SKU 补货", "先看有没有供给"],
            ["机制-加预算", "满减只投 10 天断档且核销正常", "机制对，补贴不够"],
            ["机制-停投", "OI 差的机制不要在该店继续投", "机制不对"],
            ["价格", "SKU 售价明显高于同商圈", "动销弱时先看价格"],
            ["季节+标签", "社区店加补货加活动；学校旁饮料店暑假减量", "标签 × 日历"],
        ],
    )

    add_heading(doc, "7. 总部后台", 1)
    add_body(doc, "后台回答一线今天去了哪些店、生成了哪些建议、采纳了哪些。不是经营驾驶舱，也不在本期拉通客户门店系统。")
    add_figure(doc, "04-guide-flow.png", "图 5  导购店内流程（后续场景，非本轮主路径）")

    add_heading(doc, "8. 本期范围", 1)
    add_plain(doc, "做：", bold=True, space_after=4)
    add_bullet(doc, "拍照识店、历史门店、单店销售看板")
    add_bullet(doc, "生成 AI 建议、采纳/忽略、历史记录")
    add_bullet(doc, "总部使用监控（样例 + 本机会话）")
    add_plain(doc, "不做：", bold=True, space_after=4)
    add_bullet(doc, "对接客户门店后台的上架、订货、发券实扣")
    add_bullet(doc, "客户自有门店主数据拉通")
    add_bullet(doc, "总部全国经营驾驶舱")

    add_heading(doc, "9. 演示脚本", 1)
    add_bullet(doc, "DSR 拍翡翠湾店 → 看板看 30 天三平台与机制断档 → 生成建议 → 采纳/忽略。")
    add_bullet(doc, "切换总部后台，看到本会话报告与采纳数。")
    add_bullet(doc, "可选：历史门店进望京店，建议先补供给再谈补贴。")

    add_heading(doc, "10. 附录", 1)
    add_bullet(doc, "进店先看销售看板，建议是生成出来的。")
    add_bullet(doc, "判断顺序：有没有供给，动销好不好，机制对不对，补贴够不够。")
    add_bullet(doc, "一线动作是采纳或忽略；上架做不到客户后台。")
    add_bullet(doc, "门店用 StoreVerse 的，本期不跟客户门店拉通。")

    add_heading(doc, "附图清单", 1)
    add_table(
        doc,
        ["编号", "图名", "用途"],
        [
            ["图 1", "产品能力架构", "看板到生成建议再到总部监控"],
            ["图 2", "端到端时序", "DSR 与建议引擎、总部的交互"],
            ["图 3", "数据架构", "三平台、机制、标签如何进入建议"],
            ["图 4", "DSR 主流程", "识店、看板、生成、采纳"],
            ["图 5", "导购流程", "后续场景，非本轮主讲"],
        ],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
